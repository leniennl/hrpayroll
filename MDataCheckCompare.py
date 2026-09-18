import sys
import re
from pathlib import Path
from datetime import datetime, date

import pandas as pd
import numpy as np
from openpyxl import load_workbook


# ============================================================
# SETTINGS
# ============================================================

SHEET_NAME = "MData Check"

HEADER_ROW = 3
DATA_START_ROW = 4

KEY_COLUMN = "Employee ID"
PAY_SCHEDULE_COLUMN = "PaySchedule"


# ============================================================
# FUNCTIONS
# ============================================================

def extract_version(file_path):
    """
    Extract a version such as v1.0, v1.1, v2.0 from the filename.

    Example:
        Fortnightly Employees Check PPE 20260823 v1.1.xlsm

    returns:
        v1.1
    """

    filename = Path(file_path).name

    match = re.search(r"\bv(\d+(?:\.\d+)*)\b", filename, re.IGNORECASE)

    if not match:
        # Fall back to the filename if no version is found.
        return Path(filename).stem

    return "v" + match.group(1)


def make_unique_headers(headers):
    """
    Make worksheet headers safe for use as dataframe column names.

    Blank headers are given generated names.
    Duplicate headers receive .2, .3, etc.
    """

    result = []
    counts = {}

    for position, header in enumerate(headers, start=1):

        if header is None or pd.isna(header) or str(header).strip() == "":
            header = f"Unnamed_{position}"
        else:
            header = str(header).strip()

        if header in counts:
            counts[header] += 1
            header = f"{header}.{counts[header]}"
        else:
            counts[header] = 1

        result.append(header)

    return result


def find_data_bounds(ws):
    """
    Dynamically determine the data range.

    - Header is row 3.
    - Data starts at row 4.
    - Columns extend to the last non-blank header in row 3.
    - Rows extend while column A contains an Employee ID.
    - 'Employee Total:' is treated as the end of the employee data.

    This avoids hard-coding DQ275 or any other future row/column size.
    """

    # Find last column with a header in row 3.
    last_column = 0

    for cell in ws[HEADER_ROW]:
        value = cell.value

        if value is not None and str(value).strip() != "":
            last_column = cell.column

    if last_column == 0:
        raise ValueError(
            f"No headers were found in row {HEADER_ROW} "
            f"of worksheet '{SHEET_NAME}'."
        )

    # Find the end of the employee data using column A.
    last_data_row = DATA_START_ROW - 1

    for row in range(DATA_START_ROW, ws.max_row + 1):

        employee_id = ws.cell(row=row, column=1).value

        if employee_id is None or str(employee_id).strip() == "":
            break

        if str(employee_id).strip().lower() == "employee total:":
            break

        last_data_row = row

    if last_data_row < DATA_START_ROW:
        raise ValueError(
            f"No employee data was found starting at "
            f"A{DATA_START_ROW}."
        )

    return last_data_row, last_column


def load_mdata_file(file_path):
    """
    Load only the 'MData Check' worksheet using cached formula values.

    data_only=True is intentional because this workbook contains many
    formulas and VBA. We want the displayed/calculated values rather
    than the formula expressions themselves.
    """

    print(f"Reading: {file_path.name}")

    wb = load_workbook(
        filename=file_path,
        data_only=True,
        read_only=True,
        keep_vba=True
    )

    if SHEET_NAME not in wb.sheetnames:
        wb.close()
        raise ValueError(
            f"Worksheet '{SHEET_NAME}' was not found in:\n"
            f"{file_path.name}"
        )

    ws = wb[SHEET_NAME]

    last_data_row, last_column = find_data_bounds(ws)

    # Read header row.
    raw_headers = [
        ws.cell(row=HEADER_ROW, column=col).value
        for col in range(1, last_column + 1)
    ]

    headers = make_unique_headers(raw_headers)

    # Read only the actual employee data.
    data = ws.iter_rows(
        min_row=DATA_START_ROW,
        max_row=last_data_row,
        min_col=1,
        max_col=last_column,
        values_only=True
    )

    df = pd.DataFrame(data, columns=headers)

    wb.close()

    # Required columns.
    if KEY_COLUMN not in df.columns:
        raise ValueError(
            f"'{KEY_COLUMN}' was not found in {file_path.name}."
        )

    if PAY_SCHEDULE_COLUMN not in df.columns:
        raise ValueError(
            f"'{PAY_SCHEDULE_COLUMN}' was not found in {file_path.name}."
        )

    # Employee ID must be unique.
    if df[KEY_COLUMN].isna().any():
        raise ValueError(
            f"{file_path.name} contains blank Employee ID values."
        )

    duplicates = df[df[KEY_COLUMN].duplicated(keep=False)][KEY_COLUMN].unique()

    if len(duplicates) > 0:
        raise ValueError(
            f"{file_path.name} contains duplicate Employee ID values: "
            f"{duplicates.tolist()}"
        )

    return df


def values_are_equal(old_value, new_value):
    """
    Compare two cell values while treating different types of blank
    values as equal.
    """

    old_blank = pd.isna(old_value)
    new_blank = pd.isna(new_value)

    if old_blank and new_blank:
        return True

    if old_blank != new_blank:
        return False

    # Handle datetime/date values consistently.
    if isinstance(old_value, (datetime, date)) and isinstance(
        new_value, (datetime, date)
    ):
        return old_value == new_value

    # Numeric values such as 1 and 1.0 should be considered equal.
    if isinstance(old_value, (int, float, np.number)) and isinstance(
        new_value, (int, float, np.number)
    ):
        try:
            return bool(np.isclose(old_value, new_value, equal_nan=True))
        except (TypeError, ValueError):
            pass

    return old_value == new_value


def compare_mdata_files(df1, df2, version1, version2):
    """
    Compare two MData Check dataframes using Employee ID.

    No old/new assumption is made.

    Returns:
        all differences
        changed cells
        added employees relative to version1
        removed employees relative to version1
    """

    ids1 = set(df1[KEY_COLUMN])
    ids2 = set(df2[KEY_COLUMN])

    added_ids = ids2 - ids1
    removed_ids = ids1 - ids2
    common_ids = ids1 & ids2

    # Columns common to both versions.
    common_columns = df1.columns.intersection(df2.columns)

    # Keep Employee ID out of the cell-by-cell comparison because it is
    # the key used to match records.
    compare_columns = [
        col for col in common_columns
        if col != KEY_COLUMN
    ]

    df1_indexed = df1.set_index(KEY_COLUMN)
    df2_indexed = df2.set_index(KEY_COLUMN)

    # --------------------------------------------------------
    # Changed cells
    # --------------------------------------------------------

    changed_records = []

    for employee_id in common_ids:

        row1 = df1_indexed.loc[employee_id]
        row2 = df2_indexed.loc[employee_id]

        # Use the employee's PaySchedule from version 2/current file
        # for filtering the report.
        pay_schedule = row2[PAY_SCHEDULE_COLUMN]

        for column in compare_columns:

            value1 = row1[column]
            value2 = row2[column]

            if values_are_equal(value1, value2):
                continue

            changed_records.append({
                "Change Type": "Changed",
                "Employee ID": employee_id,
                "PaySchedule": pay_schedule,
                "Column": column,
                version1: value1,
                version2: value2
            })

    changed_columns = [
        "Change Type",
        "Employee ID",
        "PaySchedule",
        "Column",
        version1,
        version2
    ]

    df_changed = pd.DataFrame(
        changed_records,
        columns=changed_columns
    )

    # --------------------------------------------------------
    # Added employees
    # --------------------------------------------------------

    added_records = []

    for employee_id in added_ids:

        row = df2_indexed.loc[employee_id]

        added_records.append({
            "Change Type": "Added",
            "Employee ID": employee_id,
            "PaySchedule": row[PAY_SCHEDULE_COLUMN],
            "Column": "",
            version1: np.nan,
            version2: "New Employee"
        })

    df_added = pd.DataFrame(
        added_records,
        columns=changed_columns
    )

    # --------------------------------------------------------
    # Removed employees
    # --------------------------------------------------------

    removed_records = []

    for employee_id in removed_ids:

        row = df1_indexed.loc[employee_id]

        removed_records.append({
            "Change Type": "Removed",
            "Employee ID": employee_id,
            "PaySchedule": row[PAY_SCHEDULE_COLUMN],
            "Column": "",
            version1: "Existing Employee",
            version2: np.nan
        })

    df_removed = pd.DataFrame(
        removed_records,
        columns=changed_columns
    )

    # --------------------------------------------------------
    # All differences
    # --------------------------------------------------------

    df_all = pd.concat(
        [
            df_added,
            df_removed,
            df_changed
        ],
        ignore_index=True
    )

    if not df_all.empty:
        df_all = (
            df_all
            .sort_values(
                ["Employee ID", "Column"],
                na_position="first"
            )
            .reset_index(drop=True)
        )

    return df_all, df_changed, df_added, df_removed


def format_output(writer, sheet_name, dataframe):
    """
    Write dataframe to Excel with basic usability formatting.
    """

    dataframe.to_excel(
        writer,
        sheet_name=sheet_name,
        index=False
    )

    worksheet = writer.book[sheet_name]

    # Freeze header row.
    worksheet.freeze_panes = "A2"

    # Add filter.
    if len(dataframe.columns) > 0 and len(dataframe) >= 1:
        worksheet.auto_filter.ref = worksheet.dimensions

    # Set reasonable column widths.
    for column_cells in worksheet.columns:

        max_length = 0
        column_letter = column_cells[0].column_letter

        for cell in column_cells:

            if cell.value is not None:
                max_length = max(
                    max_length,
                    len(str(cell.value))
                )

        worksheet.column_dimensions[column_letter].width = min(
            max(max_length + 2, 10),
            40
        )


def main():

    # ========================================================
    # Command-line arguments
    # ========================================================

    if len(sys.argv) != 3:

        print()
        print("=" * 65)
        print("MData Check comparison")
        print("=" * 65)
        print()
        print("Drag TWO versions of the workbook onto your BAT file.")
        print()
        print("Or run manually:")
        print(
            'py MDataCheckCompare.py "FILE1.xlsm" "FILE2.xlsm"'
        )
        print()
        sys.exit(1)

    file1 = Path(sys.argv[1])
    file2 = Path(sys.argv[2])

    # ========================================================
    # Validate files
    # ========================================================

    if not file1.exists():
        print(f"ERROR: File not found:\n{file1}")
        sys.exit(1)

    if not file2.exists():
        print(f"ERROR: File not found:\n{file2}")
        sys.exit(1)

    if file1.suffix.lower() not in [".xlsx", ".xlsm"]:
        print(f"ERROR: Unsupported file type:\n{file1.name}")
        sys.exit(1)

    if file2.suffix.lower() not in [".xlsx", ".xlsm"]:
        print(f"ERROR: Unsupported file type:\n{file2.name}")
        sys.exit(1)

    # ========================================================
    # Extract version labels
    # ========================================================

    version1 = extract_version(file1)
    version2 = extract_version(file2)

    if version1 == version2:
        print()
        print("WARNING: Both filenames have the same version label:")
        print(f"  {version1}")
        print()
        print("The comparison will still run, but the output columns")
        print("will have the same version label.")
        print()

    # ========================================================
    # Display files
    # ========================================================

    print()
    print("=" * 65)
    print("MData Check VERSION COMPARISON")
    print("=" * 65)
    print()
    print(f"File 1 : {file1.name}")
    print(f"Version: {version1}")
    print()
    print(f"File 2 : {file2.name}")
    print(f"Version: {version2}")
    print()
    print("Worksheet:", SHEET_NAME)
    print("Formula handling: cached values only (data_only=True)")
    print()

    # ========================================================
    # Load
    # ========================================================

    df1 = load_mdata_file(file1)
    df2 = load_mdata_file(file2)

    print(
        f"Version {version1}: "
        f"{len(df1):,} employees x {len(df1.columns):,} columns"
    )

    print(
        f"Version {version2}: "
        f"{len(df2):,} employees x {len(df2.columns):,} columns"
    )

    # ========================================================
    # Compare
    # ========================================================

    print()
    print("Comparing values...")

    (
        df_all,
        df_changed,
        df_added,
        df_removed
    ) = compare_mdata_files(
        df1,
        df2,
        version1,
        version2
    )

    # ========================================================
    # Output filename
    # ========================================================

    timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")

    output_folder = file2.parent

    output_file = output_folder / (
        f"MData_Check_Comparison_"
        f"{version1}_vs_{version2}_"
        f"{timestamp}.xlsx"
    )

    # ========================================================
    # Write Excel
    # ========================================================

    print("Creating Excel report...")

    with pd.ExcelWriter(
        output_file,
        engine="openpyxl"
    ) as writer:

        format_output(
            writer,
            "All Differences",
            df_all
        )

        format_output(
            writer,
            "Changed Cells",
            df_changed
        )

        format_output(
            writer,
            "Added Employees",
            df_added
        )

        format_output(
            writer,
            "Removed Employees",
            df_removed
        )

    # ========================================================
    # Summary
    # ========================================================

    print()
    print("=" * 65)
    print("COMPARISON COMPLETE")
    print("=" * 65)
    print()
    print(f"Version {version1} employees : {len(df1):,}")
    print(f"Version {version2} employees : {len(df2):,}")
    print(f"Added employees             : {len(df_added):,}")
    print(f"Removed employees           : {len(df_removed):,}")
    print(f"Changed cells               : {len(df_changed):,}")
    print(
        f"Employees affected         : "
        f"{df_all[KEY_COLUMN].nunique() if not df_all.empty else 0:,}"
    )
    print()
    print("Output file:")
    print(output_file)
    print()
    print("=" * 65)


if __name__ == "__main__":
    main()
