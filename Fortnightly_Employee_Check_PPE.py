"""
Fortnightly Employee Check PPE — standalone script

Usage:
    python "Fortnightly_Employee_Check_PPE.py" "C:\path\input.xlsm"

Designed for drag-and-drop through a .bat file.
The first command-line argument is the input workbook.

Outputs:
    <input name> - Issues.xlsx
    <input name> - Verification.xlsx
"""

from pathlib import Path
import re
import sys
import traceback

import numpy as np
import pandas as pd
from openpyxl.utils import get_column_letter


# ============================================================
# 0. INPUT / SETTINGS
# ============================================================

if len(sys.argv) < 2:
    raise SystemExit(
        "\nNo input file was supplied.\n"
        "Drag an Excel workbook onto the .bat file, or run:\n"
        'python "Fortnightly_Employee_Check_PPE.py" "C:\\path\\file.xlsm"\n'
    )

INPUT_FILE = Path(sys.argv[1]).resolve()

TIMESHEET_SHEET = "Timesheets Award Interpretation"
LEAVE_SHEET = "EH Leave Taken"
MANAGER_SHEET = "ManagerMapping"
EMPLOYEELISTING_SHEET = "EH Employee List"

pd.set_option("display.max_columns", 50)
pd.set_option("display.width", 180)
pd.set_option("display.max_colwidth", 60)


# ============================================================
# DIAGNOSTIC HELPERS
# ============================================================

CURRENT_STEP = "Initialisation"


def step(title):
    global CURRENT_STEP
    CURRENT_STEP = title
    print("\n" + "=" * 80)
    print(title)
    print("=" * 80)


def preview(df, rows=10, name=None):
    """Console replacement for Jupyter display()."""
    if name:
        print(f"\n{name}:")
    if isinstance(df, pd.DataFrame):
        if df.empty:
            print("<empty DataFrame>")
        else:
            print(df.head(rows).to_string(index=False))
    else:
        print(df)


def clean_columns(columns):
    return [str(c).strip() for c in columns]


def autofit_worksheet(ws, max_width=60):
    """Auto-fit columns, capped so very long text does not create huge columns."""
    for column_cells in ws.columns:
        max_length = 0
        column_letter = get_column_letter(column_cells[0].column)

        for cell in column_cells:
            if cell.value is not None:
                max_length = max(max_length, len(str(cell.value)))

        ws.column_dimensions[column_letter].width = min(max_length + 2, max_width)


def export_dataframe_excel(df, output_file, sheet_name):
    """Export one DataFrame and auto-fit the worksheet."""
    with pd.ExcelWriter(output_file, engine="openpyxl") as writer:
        df.to_excel(writer, sheet_name=sheet_name, index=False)
        ws = writer.sheets[sheet_name]

        headers = {
            cell.value: cell.column
            for cell in ws[1]
        }

        if "Date" in headers:
            date_col = headers["Date"]
            for row in ws.iter_rows(
                min_col=date_col,
                max_col=date_col,
                min_row=2,
            ):
                cell = row[0]
                if cell.value is not None:
                    cell.number_format = "d/mm/yyyy"

        # Keep Employee ID numeric in the exported workbook.
        if "Employee ID" in headers:
            employee_id_col = headers["Employee ID"]
            for row in ws.iter_rows(
                min_col=employee_id_col,
                max_col=employee_id_col,
                min_row=2,
            ):
                cell = row[0]
                if cell.value is not None:
                    cell.number_format = "0"

        autofit_worksheet(ws)


# ============================================================
# MAIN PROGRAM
# ============================================================

try:
    step("0. Input file")

    print("Input:", INPUT_FILE)
    print("Exists:", INPUT_FILE.exists())

    if not INPUT_FILE.exists():
        raise FileNotFoundError(f"Input file does not exist: {INPUT_FILE}")

    if INPUT_FILE.suffix.lower() not in {".xlsx", ".xlsm"}:
        raise ValueError(
            f"Unsupported input file type: {INPUT_FILE.suffix}. "
            "Expected .xlsx or .xlsm."
        )


    # ========================================================
    # 1. READ 4 SHEETS
    # ========================================================

    step("1. Read 4 source sheets")

    # header=1 means Excel row 2.
    df_timesheet_raw = pd.read_excel(
        INPUT_FILE,
        sheet_name=TIMESHEET_SHEET,
        header=1,
        engine="openpyxl",
    )

    df_leave_raw = pd.read_excel(
        INPUT_FILE,
        sheet_name=LEAVE_SHEET,
        header=0,
        engine="openpyxl",
    )

    df_manager_raw = pd.read_excel(
        INPUT_FILE,
        sheet_name=MANAGER_SHEET,
        header=0,
        engine="openpyxl",
    )

    df_employeelisting_raw = pd.read_excel(
        INPUT_FILE,
        sheet_name=EMPLOYEELISTING_SHEET,
        header=0,
        engine="openpyxl",
    )

    print("Timesheet:", df_timesheet_raw.shape)
    print("Leave:", df_leave_raw.shape)
    print("Manager Mapping:", df_manager_raw.shape)
    print("Employee Listing:", df_employeelisting_raw.shape)

    preview(df_timesheet_raw, name="Timesheet")
    preview(df_leave_raw, name="Leave")
    preview(df_manager_raw, name="Manager Mapping")
    preview(df_employeelisting_raw, name="Employee Listing")


    # ========================================================
    # 2. INSPECT EXACT HEADERS
    # ========================================================

    step("2. Inspect exact source headers")

    print("TIMESHEET COLUMNS")
    for i, c in enumerate(df_timesheet_raw.columns):
        print(i, repr(c))

    print("\nLEAVE COLUMNS")
    for i, c in enumerate(df_leave_raw.columns):
        print(i, repr(c))

    print("\nMANAGER MAPPING COLUMNS")
    for i, c in enumerate(df_manager_raw.columns):
        print(i, repr(c))

    print("\nEMPLOYEE LISTING COLUMNS")
    for i, c in enumerate(df_employeelisting_raw.columns):
        print(i, repr(c))


    # ========================================================
    # 3. CREATE df_timesheet
    # ========================================================

    step("3. Create df_timesheet")

    timesheet_columns = [
        "Payroll ID",
        "Name",
        "Shift Start Day & Time",
        "Unpaid breaks",
        "Timesheet Duration with break",
        "Shift Duration W/O Break",
    ]

    missing = [
        c for c in timesheet_columns
        if c not in df_timesheet_raw.columns
    ]

    if missing:
        raise KeyError(f"Missing Timesheet columns: {missing}")

    df_timesheet = df_timesheet_raw[timesheet_columns].copy()

    df_timesheet["Payroll ID"] = (
        df_timesheet["Payroll ID"]
        .astype("string")
        .str.strip()
    )

    df_timesheet["Name"] = (
        df_timesheet["Name"]
        .astype("string")
        .str.strip()
    )

    df_timesheet["Shift Start Day & Time"] = pd.to_datetime(
        df_timesheet["Shift Start Day & Time"],
        errors="coerce",
        dayfirst=True,
    ).dt.normalize()

    for col in [
        "Unpaid breaks",
        "Timesheet Duration with break",
        "Shift Duration W/O Break",
    ]:
        df_timesheet[col] = pd.to_numeric(
            df_timesheet[col],
            errors="coerce",
        )

    # Separate date key used for joins/checks.
    df_timesheet["Date"] = df_timesheet["Shift Start Day & Time"]

    preview(df_timesheet, 10, "df_timesheet")
    print("\ndf_timesheet dtypes:")
    print(df_timesheet.dtypes)


    # ========================================================
    # 4. SELECT LEAVE COLUMNS
    # ========================================================

    step("4. Select EH Leave Taken columns")

    leave_columns = [
        "EMPID",
        "Personnel",
        "Leave Category",
        "Leave Hours True in this PP",
        "Leave Dates True in this PP",
    ]

    missing = [
        c for c in leave_columns
        if c not in df_leave_raw.columns
    ]

    if missing:
        raise KeyError(f"Missing Leave columns: {missing}")

    df_leave_source = df_leave_raw[leave_columns].copy()

    preview(df_leave_source, 10, "df_leave_source")


    # ========================================================
    # 5. LEAVE PARSING / ALLOCATION HELPERS
    # ========================================================

    step("5. Define leave parsing/allocation helpers")

    def parse_leave_dates(value):
        if pd.isna(value):
            return []

        text = str(value).strip()

        if not text:
            return []

        parts = [
            p.strip()
            for p in text.split(",")
            if p.strip()
        ]

        dates = []

        for part in parts:
            dt = pd.to_datetime(
                part,
                errors="coerce",
                dayfirst=True,
            )

            if pd.notna(dt):
                dates.append(dt.normalize())

        return dates


    def allocate_leave_hours(total_hours, dates, daily_cap=7.6):
        if pd.isna(total_hours):
            total_hours = 0.0

        total_hours = float(total_hours)

        if not dates:
            return []

        # Preserve the actual value for a single-date entry,
        # e.g. 8.0.
        if len(dates) == 1:
            return [(dates[0], total_hours)]

        remaining = total_hours
        allocated = []

        for dt in dates:
            hours = min(
                daily_cap,
                max(remaining, 0),
            )

            allocated.append((dt, hours))
            remaining -= hours

        if remaining > 0:
            print(
                f"WARNING: {remaining:.2f} leave hours could not be "
                f"allocated across {len(dates)} supplied dates."
            )

        return allocated


    # ========================================================
    # 6. BUILD EXPANDED df_leave
    # ========================================================

    step("6. Build expanded df_leave")

    leave_rows = []

    for _, row in df_leave_source.iterrows():
        dates = parse_leave_dates(
            row["Leave Dates True in this PP"]
        )

        total_hours = pd.to_numeric(
            row["Leave Hours True in this PP"],
            errors="coerce",
        )

        for leave_date, daily_hours in allocate_leave_hours(
            total_hours,
            dates,
        ):
            leave_rows.append({
                "EMPID": row["EMPID"],
                "Personnel": row["Personnel"],
                "Leave Category": row["Leave Category"],
                "Leave Hours True in this PP": daily_hours,
                "Leave Dates True in this PP": leave_date,
            })

    df_leave = pd.DataFrame(leave_rows)

    preview(df_leave, 20, "df_leave")
    print("df_leave shape:", df_leave.shape)


    # ========================================================
    # 7. STANDARDISE df_leave
    # ========================================================

    step("7. Standardise Employee IDs and leave data")

    def standardise_employee_id(value):
        if pd.isna(value):
            return pd.NA

        text = str(value).strip()

        # Excel may turn integer-looking IDs into e.g. "12345.0".
        if re.fullmatch(r"\d+\.0+", text):
            text = text.split(".")[0]

        return text


    df_timesheet["Payroll ID"] = (
        df_timesheet["Payroll ID"]
        .map(standardise_employee_id)
    )

    df_leave["EMPID"] = (
        df_leave["EMPID"]
        .map(standardise_employee_id)
    )

    df_leave["Personnel"] = (
        df_leave["Personnel"]
        .astype("string")
        .str.strip()
    )

    df_leave["Leave Dates True in this PP"] = pd.to_datetime(
        df_leave["Leave Dates True in this PP"],
        errors="coerce",
    ).dt.normalize()

    df_leave["Leave Hours True in this PP"] = pd.to_numeric(
        df_leave["Leave Hours True in this PP"],
        errors="coerce",
    )

    preview(df_leave, 20, "df_leave")
    print("\ndf_leave dtypes:")
    print(df_leave.dtypes)


    # ========================================================
    # 8. DETECT FORTNIGHT
    # ========================================================

    step("8. Detect fortnight")

    all_dates = pd.concat(
        [
            df_timesheet["Date"].dropna(),
            df_leave[
                "Leave Dates True in this PP"
            ].dropna(),
        ],
        ignore_index=True,
    )

    if all_dates.empty:
        raise ValueError("No valid dates found.")

    min_date = all_dates.min()
    max_date = all_dates.max()

    fortnight_start = (
        min_date
        - pd.Timedelta(days=min_date.weekday())
    )

    fortnight_end = (
        fortnight_start
        + pd.Timedelta(days=13)
    )

    print("Minimum source date:", min_date.date())
    print("Maximum source date:", max_date.date())
    print("Fortnight start:", fortnight_start.date())
    print("Fortnight end:", fortnight_end.date())

    if max_date > fortnight_end:
        raise ValueError(
            "Source data extends beyond the detected Monday-Sunday "
            "fortnight. The workbook may contain more than one "
            "pay period."
        )


    # ========================================================
    # 9. DISPLAY FORTNIGHT CALENDAR
    # ========================================================

    step("9. Build fortnight calendar")

    fortnight_dates = pd.date_range(
        fortnight_start,
        fortnight_end,
        freq="D",
    )

    fortnight_calendar = pd.DataFrame({
        "Date": fortnight_dates
    })

    fortnight_calendar["Day"] = (
        fortnight_calendar["Date"]
        .dt.day_name()
    )

    fortnight_calendar["Weekday"] = (
        fortnight_calendar["Date"]
        .dt.weekday < 5
    )

    preview(
        fortnight_calendar,
        20,
        "fortnight_calendar",
    )


    # ========================================================
    # 10. AGGREGATE TIMESHEETS TO EMPLOYEE/DATE
    # ========================================================

    step("10. Aggregate timesheets to employee/date")

    df_ts_daily = (
        df_timesheet
        .groupby(
            [
                "Payroll ID",
                "Name",
                "Date",
            ],
            as_index=False,
            dropna=False,
        )
        .agg({
            "Timesheet Duration with break": "sum",
            "Unpaid breaks": "sum",
            "Shift Duration W/O Break": "sum",
        })
    )

    preview(df_ts_daily, 20, "df_ts_daily")


    # ========================================================
    # 11. AGGREGATE LEAVE BY EMPLOYEE/DATE/CATEGORY
    # ========================================================

    step("11. Aggregate leave by employee/date/category")

    df_leave_daily = (
        df_leave
        .groupby(
            [
                "EMPID",
                "Personnel",
                "Leave Dates True in this PP",
                "Leave Category",
            ],
            as_index=False,
            dropna=False,
        )["Leave Hours True in this PP"]
        .sum()
    )

    preview(df_leave_daily, 20, "df_leave_daily")


    # ========================================================
    # 12. COLLAPSE LEAVE CATEGORIES FOR df_master
    # ========================================================

    step("12. Build df_leave_master")

    df_leave_master = (
        df_leave_daily
        .groupby(
            [
                "EMPID",
                "Personnel",
                "Leave Dates True in this PP",
            ],
            as_index=False,
            dropna=False,
        )
        .agg({
            "Leave Category": lambda s:
                " | ".join(
                    sorted(
                        set(
                            str(x)
                            for x in s.dropna()
                        )
                    )
                ),
            "Leave Hours True in this PP": "sum",
        })
        .rename(columns={
            "EMPID": "Employee ID",
            "Personnel": "Employee Name",
            "Leave Dates True in this PP": "Date",
        })
    )

    preview(
        df_leave_master,
        20,
        "df_leave_master",
    )


    # ========================================================
    # 13. BUILD df_master
    # ========================================================

    step("13. Build df_master")

    df_ts_master = df_ts_daily.rename(columns={
        "Payroll ID": "Employee ID",
        "Name": "Employee Name",
    })

    df_ts_master = df_ts_master[
        [
            "Employee ID",
            "Employee Name",
            "Date",
            "Timesheet Duration with break",
            "Unpaid breaks",
            "Shift Duration W/O Break",
        ]
    ]

    df_leave_master = df_leave_master[
        [
            "Employee ID",
            "Employee Name",
            "Date",
            "Leave Category",
            "Leave Hours True in this PP",
        ]
    ]

    df_master = pd.merge(
        df_ts_master,
        df_leave_master,
        on=["Employee ID", "Date"],
        how="outer",
        suffixes=("_TS", "_Leave"),
    )

    df_master["Employee Name"] = (
        df_master["Employee Name_TS"]
        .combine_first(
            df_master["Employee Name_Leave"]
        )
    )

    df_master = df_master[
        [
            "Employee ID",
            "Employee Name",
            "Date",
            "Timesheet Duration with break",
            "Unpaid breaks",
            "Shift Duration W/O Break",
            "Leave Category",
            "Leave Hours True in this PP",
        ]
    ].sort_values(
        ["Employee ID", "Date"]
    ).reset_index(drop=True)

    for col in [
        "Timesheet Duration with break",
        "Unpaid breaks",
        "Shift Duration W/O Break",
        "Leave Hours True in this PP",
    ]:
        df_master[col] = pd.to_numeric(
            df_master[col],
            errors="coerce",
        )

    df_master["Date"] = pd.to_datetime(
        df_master["Date"],
        errors="coerce",
    ).dt.normalize()

    preview(df_master, 30, "df_master")
    print("df_master shape:", df_master.shape)


    # ========================================================
    # 14. ISSUE A
    # ========================================================

    step("14. Issue A — No Unpaid Break in an Over 5h Shift")

    mask_a = (
        (df_master["Timesheet Duration with break"] > 5)
        & (
            df_master["Unpaid breaks"].isna()
            | (df_master["Unpaid breaks"] <= 0)
        )
    )

    issues_a = df_master.loc[
        mask_a,
        [
            "Employee ID",
            "Employee Name",
            "Date",
            "Timesheet Duration with break",
            "Unpaid breaks",
            "Shift Duration W/O Break",
            "Leave Category",
            "Leave Hours True in this PP",
        ],
    ].copy()

    issues_a["Issue"] = (
        "No Unpaid Break in an Over 5h Shift"
    )

    print("Issue A count:", len(issues_a))
    preview(issues_a, 50, "issues_a")


    # ========================================================
    # 15. ISSUE B
    # ========================================================

    step("15. Issue B — Unpaid Break in an Under 5h Shift")

    mask_b = (
        (df_master["Timesheet Duration with break"] < 5)
        & (df_master["Unpaid breaks"] > 0)
    )

    issues_b = df_master.loc[
        mask_b,
        [
            "Employee ID",
            "Employee Name",
            "Date",
            "Timesheet Duration with break",
            "Unpaid breaks",
            "Shift Duration W/O Break",
            "Leave Category",
            "Leave Hours True in this PP",
        ],
    ].copy()

    issues_b["Issue"] = (
        "Unpaid Break in an Under 5h Shift"
    )

    print("Issue B count:", len(issues_b))
    preview(issues_b, 50, "issues_b")


    # ========================================================
    # 16. ISSUE C
    # ========================================================

    step("16. Issue C — Leave Over Applied")

    mask_c = (
        df_master["Shift Duration W/O Break"].notna()
        & df_master["Leave Hours True in this PP"].notna()
        & (
            df_master["Shift Duration W/O Break"].fillna(0)
            + df_master["Leave Hours True in this PP"].fillna(0)
            > 7.61
        )
    )

    issues_c = df_master.loc[
        mask_c,
        [
            "Employee ID",
            "Employee Name",
            "Date",
            "Timesheet Duration with break",
            "Unpaid breaks",
            "Shift Duration W/O Break",
            "Leave Category",
            "Leave Hours True in this PP",
        ],
    ].copy()

    issues_c["Issue"] = "Leave Over Applied"

    print("Issue C count:", len(issues_c))
    preview(issues_c, 50, "issues_c")


    # ========================================================
    # 17. EMPLOYEE POPULATION
    # ========================================================

    step("17. Build employee population")

    employee_population = df_employeelisting_raw[
        df_employeelisting_raw["PaySchedule"]
        == "Walkinshaw Automotive Fortnightly"
    ][
        ["EmployeeId", "EMPFullName"]
    ].rename(columns={
        "EmployeeId": "Employee ID",
        "EMPFullName": "Employee Name",
    })

    employee_population["Employee ID"] = (
        employee_population["Employee ID"]
        .map(standardise_employee_id)
    )

    employee_population["Employee Name"] = (
        employee_population["Employee Name"]
        .astype("string")
        .str.strip()
    )

    employee_population = (
        employee_population
        .sort_values(
            ["Employee ID", "Employee Name"]
        )
        .drop_duplicates(
            subset=["Employee ID"],
            keep="first",
        )
        .reset_index(drop=True)
    )

    print(
        "Employees checked:",
        len(employee_population),
    )

    preview(
        employee_population,
        100,
        "employee_population",
    )


    # ========================================================
    # 18. EMPLOYEE x WEEKDAY MATRIX / ISSUE E
    # ========================================================

    step("18. Missing Leave Form & Timesheet check")

    weekdays = pd.date_range(
        fortnight_start,
        fortnight_end,
        freq="B",
    )

    employee_days = (
        employee_population
        .assign(key=1)
        .merge(
            pd.DataFrame({
                "Date": weekdays,
                "key": 1,
            }),
            on="key",
        )
        .drop(columns="key")
    )

    ts_keys = (
        df_ts_master[
            ["Employee ID", "Date"]
        ]
        .drop_duplicates()
        .assign(HasTimesheet=True)
    )

    leave_keys = (
        df_leave_master[
            ["Employee ID", "Date"]
        ]
        .drop_duplicates()
        .assign(HasLeave=True)
    )

    missing_check = (
        employee_days
        .merge(
            ts_keys,
            on=["Employee ID", "Date"],
            how="left",
        )
        .merge(
            leave_keys,
            on=["Employee ID", "Date"],
            how="left",
        )
    )

    missing_rows = missing_check[
        missing_check["HasTimesheet"].isna()
        & missing_check["HasLeave"].isna()
    ].copy()

    issues_e = missing_rows[
        [
            "Employee ID",
            "Employee Name",
            "Date",
        ]
    ].copy()

    issues_e[
        "Timesheet Duration with break"
    ] = np.nan

    issues_e["Unpaid breaks"] = np.nan
    issues_e["Shift Duration W/O Break"] = np.nan
    issues_e["Leave Category"] = pd.NA
    issues_e[
        "Leave Hours True in this PP"
    ] = np.nan

    issues_e["Issue"] = (
        "Missing Leave Form & Timesheet"
    )

    print("Issue E count:", len(issues_e))
    preview(issues_e, 100, "issues_e")


    # ========================================================
    # 19. BUILD df_issues
    # ========================================================

    step("19. Combine all issues into df_issues")

    issue_columns = [
        "Employee ID",
        "Employee Name",
        "Date",
        "Timesheet Duration with break",
        "Unpaid breaks",
        "Shift Duration W/O Break",
        "Leave Category",
        "Leave Hours True in this PP",
        "Issue",
    ]

    def prepare_issue_frame(df):
        out = df.copy()

        for col in issue_columns:
            if col not in out.columns:
                out[col] = np.nan

        return out[issue_columns]


    df_issues = pd.concat(
        [
            prepare_issue_frame(issues_a),
            prepare_issue_frame(issues_b),
            prepare_issue_frame(issues_c),
            # Issue D remains intentionally disabled, matching the finished notebook.
            prepare_issue_frame(issues_e),
        ],
        ignore_index=True,
    )

    df_issues["Employee ID"] = (
        df_issues["Employee ID"]
        .map(standardise_employee_id)
    )

    df_issues["Employee Name"] = (
        df_issues["Employee Name"]
        .astype("string")
        .str.strip()
    )

    df_issues["Date"] = pd.to_datetime(
        df_issues["Date"],
        errors="coerce",
    ).dt.normalize()

    for col in [
        "Timesheet Duration with break",
        "Unpaid breaks",
        "Shift Duration W/O Break",
        "Leave Hours True in this PP",
    ]:
        df_issues[col] = pd.to_numeric(
            df_issues[col],
            errors="coerce",
        )

    df_issues = df_issues.sort_values(
        ["Employee ID", "Date", "Issue"]
    ).reset_index(drop=True)

    print("Total issues:", len(df_issues))
    preview(df_issues, 100, "df_issues")


    # ========================================================
    # 20. MANAGER MAPPING
    # ========================================================

    step("20. Inspect ManagerMapping and create manager lookup")

    preview(
        df_manager_raw,
        10,
        "df_manager_raw",
    )

    for i, c in enumerate(df_manager_raw.columns):
        print(i, repr(c))

    EMPLOYEE_ID_COL = "EmployeeID"
    EMPLOYEE_NAME_COL = "Personnel"
    MANAGER_EMAIL_COL = "Manager Work Email"
    SECONDARY_MANAGER_EMAIL_COL = "Secondary manager Email"

    required_columns = [
        EMPLOYEE_NAME_COL,
        EMPLOYEE_ID_COL,
        MANAGER_EMAIL_COL,
        SECONDARY_MANAGER_EMAIL_COL,
    ]

    missing_columns = [
        c
        for c in required_columns
        if c not in df_manager_raw.columns
    ]

    if missing_columns:
        raise ValueError(
            "Missing required columns in df_manager_raw: "
            f"{missing_columns}"
        )

    print("Employee ID:", EMPLOYEE_ID_COL)
    print("Employee Name:", EMPLOYEE_NAME_COL)
    print("Manager Work Email:", MANAGER_EMAIL_COL)
    print(
        "Secondary Manager Email:",
        SECONDARY_MANAGER_EMAIL_COL,
    )

    df_manager = df_manager_raw[
        [
            EMPLOYEE_ID_COL,
            EMPLOYEE_NAME_COL,
            MANAGER_EMAIL_COL,
            SECONDARY_MANAGER_EMAIL_COL,
        ]
    ].copy()

    df_manager.columns = [
        "Employee ID",
        "Mapping Employee Name",
        "Manager",
        "Secondary Manager",
    ]

    df_manager["Employee ID"] = (
        df_manager["Employee ID"]
        .map(standardise_employee_id)
    )

    df_manager["Mapping Employee Name"] = (
        df_manager["Mapping Employee Name"]
        .astype("string")
        .str.strip()
    )

    df_manager["Manager"] = (
        df_manager["Manager"]
        .astype("string")
        .str.strip()
    )

    df_manager["Secondary Manager"] = (
        df_manager["Secondary Manager"]
        .astype("string")
        .str.strip()
    )

    # Convert blank strings to <NA>.
    df_manager["Manager"] = (
        df_manager["Manager"].replace("", pd.NA)
    )

    df_manager["Secondary Manager"] = (
        df_manager["Secondary Manager"]
        .replace("", pd.NA)
    )

    duplicate_manager_ids = df_manager.loc[
        df_manager["Employee ID"].duplicated(keep=False)
    ].sort_values("Employee ID")

    if not duplicate_manager_ids.empty:
        print(
            "\nWARNING: duplicate Employee IDs "
            "in ManagerMapping:"
        )
        preview(
            duplicate_manager_ids,
            100,
            "duplicate_manager_ids",
        )

    df_manager_lookup = (
        df_manager
        .drop_duplicates(
            subset=["Employee ID"],
            keep="first",
        )
        [
            [
                "Employee ID",
                "Manager",
                "Secondary Manager",
            ]
        ]
    )

    df_issues = df_issues.merge(
        df_manager_lookup,
        on="Employee ID",
        how="left",
    )

    preview(
        df_manager_lookup,
        50,
        "df_manager_lookup",
    )

    preview(
        df_issues,
        100,
        "df_issues after Manager Mapping",
    )


    # ========================================================
    # 21. VALIDATION SUMMARY
    # ========================================================

    step("21. Validation summary")

    print("========== SUMMARY ==========")
    print("Timesheet rows:", len(df_timesheet))
    print(
        "Daily timesheet employee/date rows:",
        len(df_ts_daily),
    )
    print("Expanded leave rows:", len(df_leave))
    print(
        "Daily leave employee/date rows:",
        len(df_leave_master),
    )
    print("Master rows:", len(df_master))
    print("Issue rows:", len(df_issues))
    print(
        "Employees checked:",
        employee_population["Employee ID"].nunique(),
    )

    print("\n========== ISSUES BY TYPE ==========")

    issue_summary = (
        df_issues["Issue"]
        .value_counts()
        .rename_axis("Issue")
        .reset_index(name="Count")
    )

    preview(
        issue_summary,
        100,
        "Issues by type",
    )

    print(
        "\n========== ISSUE EMPLOYEES "
        "WITHOUT MANAGER MAPPING =========="
    )

    unmapped_issue_employees = (
        df_issues.loc[
            df_issues["Manager"].isna(),
            ["Employee ID", "Employee Name"],
        ]
        .drop_duplicates()
    )

    preview(
        unmapped_issue_employees,
        100,
        "Issue employees without Manager",
    )


    # ========================================================
    # 22. CHECK ALL SOURCE EMPLOYEES FOR MANAGER MAPPING
    # ========================================================

    step("22. Check all source employees for Manager Mapping")

    source_manager_check = employee_population.merge(
        df_manager_lookup,
        on="Employee ID",
        how="left",
    )

    unmapped_source = source_manager_check.loc[
        source_manager_check["Manager"].isna()
    ]

    print(
        "Source employees without ManagerMapping:",
        len(unmapped_source),
    )

    preview(
        unmapped_source,
        100,
        "unmapped_source",
    )


    # ========================================================
    # 23. EXPORT df_issues
    # ========================================================

    step("23. Export df_issues")

    OUTPUT_FILE = INPUT_FILE.with_name(
        INPUT_FILE.stem + " - Issues.xlsx"
    )

    # Convert Employee ID to numeric for the final Excel output,
    # matching the finished notebook.
    df_issues["Employee ID"] = pd.to_numeric(
        df_issues["Employee ID"],
        errors="coerce",
    )

    export_dataframe_excel(
        df_issues,
        OUTPUT_FILE,
        "df_issues",
    )

    print("Created:", OUTPUT_FILE)
    print("Exists:", OUTPUT_FILE.exists())


    # ========================================================
    # 24. VERIFICATION EXPORT
    # ========================================================

    step("24. Create verification workbook")

    VERIFICATION_FILE = INPUT_FILE.with_name(
        INPUT_FILE.stem + " - Verification.xlsx"
    )

    with pd.ExcelWriter(
        VERIFICATION_FILE,
        engine="openpyxl",
    ) as writer:

        df_timesheet.to_excel(
            writer,
            sheet_name="df_timesheet",
            index=False,
        )

        df_leave.to_excel(
            writer,
            sheet_name="df_leave",
            index=False,
        )

        df_master.to_excel(
            writer,
            sheet_name="df_master",
            index=False,
        )

        df_issues.to_excel(
            writer,
            sheet_name="df_issues",
            index=False,
        )

        for ws in writer.book.worksheets:
            autofit_worksheet(ws)

            # Apply date formatting where a Date column exists.
            header_map = {
                cell.value: cell.column
                for cell in ws[1]
            }

            if "Date" in header_map:
                date_col = header_map["Date"]
                for row in ws.iter_rows(
                    min_col=date_col,
                    max_col=date_col,
                    min_row=2,
                ):
                    if row[0].value is not None:
                        row[0].number_format = "d/mm/yyyy"

            # Apply numeric Employee ID formatting where applicable.
            if "Employee ID" in header_map:
                employee_id_col = header_map["Employee ID"]
                for row in ws.iter_rows(
                    min_col=employee_id_col,
                    max_col=employee_id_col,
                    min_row=2,
                ):
                    if row[0].value is not None:
                        row[0].number_format = "0"

    print("Created:", VERIFICATION_FILE)
    print("Exists:", VERIFICATION_FILE.exists())


    # ========================================================
    # COMPLETE
    # ========================================================

    step("COMPLETE")

    print("Input file:")
    print(INPUT_FILE)

    print("\nIssues output:")
    print(OUTPUT_FILE)

    print("\nVerification output:")
    print(VERIFICATION_FILE)

    print("\nTotal issue rows:", len(df_issues))

    print("\nScript completed successfully.")


except Exception as exc:
    print("\n")
    print("#" * 80)
    print("ERROR — SCRIPT STOPPED")
    print("#" * 80)
    print("Step:", CURRENT_STEP)
    print("Error type:", type(exc).__name__)
    print("Error message:", str(exc))
    print("\nFull traceback:")
    traceback.print_exc()
    print("#" * 80)

    # Keep the .bat window open so the error can be read.
    if sys.stdin.isatty():
        input("\nPress Enter to close...")

    raise
